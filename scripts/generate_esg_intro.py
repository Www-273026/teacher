from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "ppt" / "ESG理论与实践_第一章_绪论.pptx"

W, H = 13.333, 7.5
NAVY = "10243A"
INK = "183149"
MUTED = "617487"
PAPER = "F5F7F8"
WHITE = "FFFFFF"
GREEN = "2C8C70"
GREEN_LIGHT = "E6F3ED"
ORANGE = "D57B46"
ORANGE_LIGHT = "FBF0E8"
BLUE = "4879B5"
BLUE_LIGHT = "EAF0F8"
GOLD = "E8B75A"
LINE = "DCE4E9"
FONT = "Microsoft YaHei"

SOURCES = [
    ("UN Global Compact, Who Cares Wins (2004)",
     "https://www.unglobalcompact.org/docs/issues_doc/Financial_markets/who_cares_who_wins.pdf"),
    ("UN News, Principles for Responsible Investment launch (2006)",
     "https://news.un.org/zh/story/2006/04/53132"),
    ("United Nations, Transforming our world: the 2030 Agenda (2015)",
     "https://sdgs.un.org/2030agenda"),
    ("IFRS Foundation, ISSB issues IFRS S1 and IFRS S2 (2023)",
     "https://www.ifrs.org/news-and-events/news/2023/06/issb-issues-ifrs-s1-ifrs-s2/"),
    ("中共中央、国务院：关于完整准确全面贯彻新发展理念做好碳达峰碳中和工作的意见（2021）",
     "https://www.gov.cn/zhengce/2021-10/24/content_5644613.htm"),
    ("国务院：2030年前碳达峰行动方案（2021）",
     "https://www.gov.cn/zhengce/content/2021-10/26/content_5644984.htm"),
    ("上海证券交易所：上市公司自律监管指引第14号——可持续发展报告（试行）（2024）",
     "https://www.sse.com.cn/lawandrules/sselawsrules/stocks/mainipo/c/c_20240412_10753180.shtml"),
]


def color(value):
    return RGBColor.from_string(value)


def set_background(slide, value=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color(value)


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    return shape


def text(slide, value, x, y, w, h, size=18, fill=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.02,
         font=FONT, spacing=1.08):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(5)
        paragraph.line_spacing = spacing
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color(fill)
    return shape


def rich_text(slide, parts, x, y, w, h, size=14, fill=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.space_after = Pt(0)
    for label, url in parts:
        run = paragraph.add_run()
        run.text = label
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color(fill)
        run.hyperlink.address = url
    return shape


def new_slide(title, section, page, prompt=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY if dark else PAPER)
    fg = WHITE if dark else INK
    muted = "AFC1CE" if dark else MUTED
    if not dark:
        text(slide, section.upper(), 0.62, 0.34, 8.5, 0.24, 10, GREEN, True)
        text(slide, title, 0.62, 0.72, 12.0, 0.55, 27, fg, True)
        box(slide, 0.62, 1.43, 0.55, 0.045, GREEN, radius=False)
    if prompt:
        text(slide, f"教学提示  {prompt}", 0.62, 7.08, 11.2, 0.22, 9, muted)
    text(slide, f"{page:02d}", 12.15, 7.06, 0.55, 0.22, 9, muted, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, label, heading, body, accent, light):
    box(slide, x, y, w, h, WHITE)
    box(slide, x, y, 0.08, h, accent, radius=False)
    text(slide, label, x + 0.25, y + 0.23, w - 0.5, 0.27, 11, accent, True)
    text(slide, heading, x + 0.25, y + 0.65, w - 0.5, 0.48, 20, INK, True)
    text(slide, body, x + 0.25, y + 1.25, w - 0.5, h - 1.45, 14, MUTED)


def add_linked_source(slide, label, url, x, y, w=11.8):
    rich_text(slide, [(label, url)], x, y, w, 0.24, 9, MUTED)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = "ESG理论与实践｜第一章 绪论"
prs.core_properties.subject = "本科课堂教学课件：ESG概念、发展历程、SDGs关联与中国实践"
prs.core_properties.author = "课程教学材料"
prs.core_properties.keywords = "ESG, SDGs, 本科教学, 绪论"

# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, NAVY)
box(slide, 9.3, -0.45, 4.65, 4.65, "183A51")
box(slide, 10.15, 0.4, 2.8, 2.8, GREEN)
box(slide, 10.85, 1.12, 1.4, 1.4, NAVY)
text(slide, "课程导论  /  CHAPTER 01", 0.9, 0.85, 6.4, 0.36, 12, "AFC1CE", True)
text(slide, "ESG", 0.85, 1.48, 5.2, 1.22, 62, WHITE, True)
text(slide, "理论与实践", 0.91, 2.76, 6.3, 0.75, 34, WHITE, True)
text(slide, "第一章  绪论", 0.94, 3.68, 5.0, 0.48, 21, "C7D7E1")
box(slide, 0.94, 4.48, 0.65, 0.05, GOLD, radius=False)
text(slide, "从环境、社会与治理，理解组织的长期价值", 0.94, 4.82, 7.1, 0.45, 18, "D7E3EA")
text(slide, "本科课程 · 教学课件", 0.94, 6.55, 4.6, 0.3, 11, "AFC1CE")
text(slide, "01", 12.15, 7.06, 0.55, 0.22, 9, "AFC1CE", align=PP_ALIGN.RIGHT)

# 2. Learning goals
slide = new_slide("本章学习路线", "导入", 2,
                  "可先请学生举例：企业利润之外，还有哪些表现会影响长期发展？")
goals = [
    ("01", "说清概念", "解释 ESG 与 E / S / G 三个维度", GREEN, GREEN_LIGHT),
    ("02", "看懂体系", "描述议题如何进入管理、披露与决策", BLUE, BLUE_LIGHT),
    ("03", "联系目标", "区分 ESG 与 SDGs，并说明二者如何衔接", ORANGE, ORANGE_LIGHT),
    ("04", "回到中国", "结合政策与企业情境讨论实践意义", "8A6CB1", "F1ECF7"),
]
for i, (n, h, b, a, light) in enumerate(goals):
    x = 0.62 + (i % 2) * 6.08
    y = 1.82 + (i // 2) * 2.05
    box(slide, x, y, 5.75, 1.72, WHITE)
    box(slide, x + 0.24, y + 0.26, 0.56, 0.56, light)
    text(slide, n, x + 0.24, y + 0.4, 0.56, 0.2, 10, a, True, align=PP_ALIGN.CENTER)
    text(slide, h, x + 1.02, y + 0.28, 4.2, 0.35, 19, INK, True)
    text(slide, b, x + 1.02, y + 0.83, 4.25, 0.53, 13, MUTED)

# 3. Definition
slide = new_slide("ESG是什么？", "第一节  概念与体系", 3,
                  "强调“分析与管理视角”，不要把 ESG 等同于环保或某一个评级分数。")
box(slide, 0.62, 1.82, 7.45, 3.98, NAVY)
text(slide, "ESG 是观察组织可持续发展相关议题的一组维度。", 1.0, 2.22, 6.65, 1.08, 27, WHITE, True)
text(slide, "帮助组织识别议题、管理影响与风险，\n并沟通行动与表现。", 1.02, 3.78, 6.35, 0.93, 18, "C7D7E1")
box(slide, 8.38, 1.82, 4.3, 1.1, GREEN_LIGHT)
text(slide, "E  环境", 8.72, 2.15, 3.5, 0.34, 19, GREEN, True)
box(slide, 8.38, 3.15, 4.3, 1.1, ORANGE_LIGHT)
text(slide, "S  社会", 8.72, 3.48, 3.5, 0.34, 19, ORANGE, True)
box(slide, 8.38, 4.48, 4.3, 1.1, BLUE_LIGHT)
text(slide, "G  治理", 8.72, 4.81, 3.5, 0.34, 19, BLUE, True)
text(slide, "具体议题会因行业、地区与组织活动而不同。", 0.68, 6.24, 11.7, 0.36, 15, MUTED)

# 4. E/S/G
slide = new_slide("三个维度，关注不同问题", "第一节  概念与体系", 4,
                  "让学生把熟悉的企业或校园组织，分别放入 E、S、G 维度试着举例。")
card(slide, 0.62, 1.82, 3.83, 4.55, "E  /  ENVIRONMENT", "环境", "能源与排放\n污染防治\n水资源与生态\n\n追问：经营活动如何影响自然环境？", GREEN, GREEN_LIGHT)
card(slide, 4.75, 1.82, 3.83, 4.55, "S  /  SOCIAL", "社会", "员工权益与安全\n消费者与产品责任\n供应链与社区\n\n追问：谁受到组织活动影响？", ORANGE, ORANGE_LIGHT)
card(slide, 8.88, 1.82, 3.83, 4.55, "G  /  GOVERNANCE", "治理", "董事会与监督\n商业伦理与合规\n风险管理与透明度\n\n追问：谁负责决策、监督与问责？", BLUE, BLUE_LIGHT)

# 5. ESG management cycle
slide = new_slide("ESG如何进入管理？", "第一节  概念与体系", 5,
                  "说明这是一种便于教学的管理流程，不是唯一或强制的标准流程。")
steps = [
    ("1", "识别", "议题与相关方", GREEN),
    ("2", "评估", "影响、风险与机会", BLUE),
    ("3", "行动", "责任、目标与资源", ORANGE),
    ("4", "跟踪", "指标、数据与复盘", "8A6CB1"),
    ("5", "沟通", "披露、对话与改进", GREEN),
]
for i, (n, head, sub, accent) in enumerate(steps):
    x = 0.62 + i * 2.57
    box(slide, x, 2.45, 2.15, 2.22, WHITE)
    box(slide, x + 0.18, 2.66, 0.49, 0.49, accent)
    text(slide, n, x + 0.18, 2.79, 0.49, 0.18, 11, WHITE, True, align=PP_ALIGN.CENTER)
    text(slide, head, x + 0.2, 3.42, 1.76, 0.38, 19, INK, True)
    text(slide, sub, x + 0.2, 3.98, 1.8, 0.42, 12, MUTED)
    if i < len(steps) - 1:
        text(slide, "→", x + 2.18, 3.24, 0.34, 0.4, 21, MUTED, True, align=PP_ALIGN.CENTER)
box(slide, 0.62, 5.35, 12.08, 0.72, GREEN_LIGHT)
text(slide, "核心：把重要议题放进决策与持续改进，而不止停留在宣传或报告。", 0.92, 5.56, 11.45, 0.28, 15, GREEN, True)

# 6. Timeline
slide = new_slide("现代 ESG 发展的若干节点", "第二节  发展历程", 6,
                  "时间线是若干代表性节点，不代表 ESG 思想从某一年才开始。")
timeline = [
    ("2004", "议题进入金融讨论", "《Who Cares Wins》倡议金融市场\n更系统关注 ESG 因素。", GREEN),
    ("2006", "投资实践推进", "联合国负责任投资原则（PRI）\n发布，倡导投资决策纳入 ESG。", BLUE),
    ("2015", "全球目标形成", "联合国通过《2030 年可持续发展议程》，\n包括 17 项 SDGs。", ORANGE),
    ("2023", "披露准则发布", "ISSB 发布 IFRS S1 与 IFRS S2，\n为可持续相关披露提供基准。", "8A6CB1"),
]
box(slide, 1.02, 3.05, 11.2, 0.045, LINE, radius=False)
for i, (year, head, body, accent) in enumerate(timeline):
    x = 0.72 + i * 3.12
    box(slide, x + 0.22, 2.8, 0.5, 0.5, accent)
    text(slide, str(i + 1), x + 0.22, 2.94, 0.5, 0.18, 10, WHITE, True, align=PP_ALIGN.CENTER)
    text(slide, year, x, 1.95, 2.65, 0.53, 26, accent, True)
    text(slide, head, x, 3.58, 2.7, 0.43, 17, INK, True)
    text(slide, body, x, 4.18, 2.72, 0.9, 12, MUTED)
add_linked_source(slide, "来源：UN Global Compact（2004）；PRI；联合国《2030年议程》（2015）；IFRS Foundation（2023）。", SOURCES[0][1], 0.68, 6.38)

# 7. Scope and framework
slide = new_slide("ESG不是一个统一的“分数”", "第一节  概念与体系", 7,
                  "提醒学生：不同评级/标准的目的、范围、指标与权重可能不同，不能只看单一排名。")
items = [
    ("议题", "组织面对什么环境、社会与治理问题？", GREEN),
    ("管理", "如何设定责任、目标并采取行动？", BLUE),
    ("披露", "哪些信息需要被记录、衡量与沟通？", ORANGE),
    ("评价", "评价者依据什么方法与数据作出判断？", "8A6CB1"),
]
for i, (head, body, accent) in enumerate(items):
    y = 1.82 + i * 1.05
    box(slide, 0.82, y, 11.72, 0.78, WHITE)
    text(slide, head, 1.12, y + 0.18, 1.55, 0.34, 17, accent, True)
    text(slide, body, 2.8, y + 0.2, 8.9, 0.33, 14, INK)
text(slide, "概念视角 ≠ 披露标准 ≠ 评级方法", 1.02, 6.28, 11.0, 0.38, 18, NAVY, True, align=PP_ALIGN.CENTER)

# 8. SDGs distinction
slide = new_slide("ESG 与 SDGs：目标和视角相连，但不是一回事", "第三节  ESG 与可持续发展目标", 8,
                  "把“希望实现什么”与“组织如何识别并管理自身议题”作对照。")
box(slide, 0.62, 1.82, 5.78, 3.72, NAVY)
text(slide, "SDGs", 0.98, 2.18, 4.95, 0.62, 31, WHITE, True)
text(slide, "面向全球共同发展的目标议程", 0.98, 2.96, 4.95, 0.42, 18, "C7D7E1", True)
text(slide, "联合国《2030年议程》\n17 项目标 · 169 项具体目标\n关注经济、社会与环境的综合发展", 0.98, 3.7, 4.95, 1.14, 15, WHITE)
box(slide, 6.92, 1.82, 5.78, 3.72, WHITE)
text(slide, "ESG", 7.28, 2.18, 4.95, 0.62, 31, GREEN, True)
text(slide, "组织分析与管理议题的视角", 7.28, 2.96, 4.95, 0.42, 18, INK, True)
text(slide, "结合组织活动与相关方\n识别议题、管理影响与风险\n支持决策、行动与信息沟通", 7.28, 3.7, 4.95, 1.14, 15, MUTED)
text(slide, "SDGs 提供共同方向       ESG 帮助组织审视自身行动       两者通过具体目标与结果衔接", 0.76, 5.98, 11.9, 0.48, 15, GREEN, True, align=PP_ALIGN.CENTER)
add_linked_source(slide, "来源：联合国《2030年可持续发展议程》（A/RES/70/1）。", SOURCES[2][1], 0.68, 6.62)

# 9. Illustrative mapping
slide = new_slide("一个行动，可以连接多个目标", "第三节  ESG 与可持续发展目标", 9,
                  "这是课堂示例映射，不是官方唯一对应关系；应依据企业活动、证据与实际结果判断。")
text(slide, "示例：食品企业改善供应链与工厂运营", 0.74, 1.82, 11.8, 0.42, 19, INK, True)
rows = [
    ("E", "工厂节能、减少食物损耗", "SDG 7 清洁能源  ·  SDG 12 负责任消费和生产  ·  SDG 13 气候行动", GREEN, GREEN_LIGHT),
    ("S", "保障员工安全、改善供应商劳动条件", "SDG 3 良好健康与福祉  ·  SDG 8 体面工作", ORANGE, ORANGE_LIGHT),
    ("G", "明确董事会监督、公开目标与进展", "SDG 12 负责任消费和生产  ·  SDG 16 和平、正义与强大机构", BLUE, BLUE_LIGHT),
]
for i, (letter, action, targets, accent, light) in enumerate(rows):
    y = 2.53 + i * 1.08
    box(slide, 0.74, y, 11.85, 0.84, WHITE)
    box(slide, 0.95, y + 0.14, 0.58, 0.56, light)
    text(slide, letter, 0.95, y + 0.28, 0.58, 0.2, 16, accent, True, align=PP_ALIGN.CENTER)
    text(slide, action, 1.83, y + 0.12, 4.28, 0.3, 14, INK, True)
    text(slide, targets, 6.22, y + 0.14, 6.0, 0.48, 12, MUTED)
text(slide, "一个 SDG 可能涉及多个维度；一个 ESG 行动也可能关联多个 SDGs。", 0.9, 6.18, 11.5, 0.38, 16, NAVY, True, align=PP_ALIGN.CENTER)

# 10. China context
slide = new_slide("ESG在中国：从发展要求走向经营实践", "第四节  中国实践意义", 10,
                  "双碳目标是国家层面的方向；企业路径、指标与转型节奏需要结合行业和实际情况。")
cards = [
    ("绿色低碳转型", "关注能源、排放、资源效率\n识别转型风险与技术机会", GREEN, GREEN_LIGHT),
    ("发展与民生", "兼顾员工安全、产品责任\n供应链与社区影响", ORANGE, ORANGE_LIGHT),
    ("治理与透明", "明确管理责任、加强风险控制\n用可核验信息回应相关方", BLUE, BLUE_LIGHT),
]
for i, (head, body, accent, light) in enumerate(cards):
    x = 0.62 + i * 4.12
    box(slide, x, 1.82, 3.83, 2.35, WHITE)
    box(slide, x + 0.24, 2.08, 0.14, 0.47, accent, radius=False)
    text(slide, head, x + 0.52, 2.07, 3.0, 0.42, 17, INK, True)
    text(slide, body, x + 0.26, 2.85, 3.26, 0.89, 14, MUTED)
box(slide, 0.62, 4.52, 12.07, 1.28, NAVY)
text(slide, "国家目标", 0.98, 4.83, 1.5, 0.3, 13, "AFC1CE", True)
text(slide, "力争 2030 年前实现碳达峰，努力争取 2060 年前实现碳中和。", 2.5, 4.77, 9.55, 0.42, 18, WHITE, True)
text(slide, "2024 年沪深交易所发布上市公司可持续发展报告指引（试行），推动相关披露规范化。", 0.77, 6.18, 11.8, 0.45, 13, MUTED)
add_linked_source(slide, "来源：中共中央、国务院（2021）；国务院《2030年前碳达峰行动方案》（2021）；上交所指引（2024）。", SOURCES[5][1], 0.68, 6.65)

# 11. Classroom scenario
slide = new_slide("情境练习：新工厂选址", "第四节  中国实践意义", 11,
                  "邀请学生不要急于下结论，先区分事实、风险、受影响者与可采取的措施。")
box(slide, 0.62, 1.82, 4.05, 4.35, NAVY)
text(slide, "案例", 0.98, 2.18, 3.2, 0.38, 14, GOLD, True)
text(slide, "一家食品企业\n计划建设新工厂。", 0.98, 2.8, 3.2, 1.25, 26, WHITE, True)
text(slide, "请用 ESG 视角做初步“体检”", 0.98, 4.82, 3.2, 0.52, 15, "C7D7E1")
diagnostics = [
    ("E", "用水、能耗、排放与周边生态", GREEN),
    ("S", "员工安全、消费者与社区沟通", ORANGE),
    ("G", "选址决策、监督责任与信息披露", BLUE),
]
for i, (letter, prompt, accent) in enumerate(diagnostics):
    y = 2.05 + i * 1.21
    box(slide, 5.12, y, 7.45, 0.92, WHITE)
    text(slide, letter, 5.45, y + 0.23, 0.55, 0.28, 19, accent, True)
    text(slide, prompt, 6.2, y + 0.24, 5.95, 0.38, 15, INK, True)
box(slide, 5.12, 5.85, 7.45, 0.55, GREEN_LIGHT)
text(slide, "每组提出 1 个优先议题 + 1 条可执行建议", 5.43, 6.01, 6.8, 0.24, 13, GREEN, True)

# 12. Discussion
slide = new_slide("小组讨论：什么才算“做得好”？", "课堂讨论", 12,
                  "建议 5 分钟讨论、每组 1 分钟汇报；鼓励学生给出证据和权衡，而非只喊口号。")
prompts = [
    ("01", "谁受影响？", "识别员工、消费者、社区、供应商与投资者", GREEN),
    ("02", "优先做什么？", "说明最重要的议题及判断依据", ORANGE),
    ("03", "如何证明？", "提出一个可观察的指标或证据", BLUE),
]
for i, (n, heading, body, accent) in enumerate(prompts):
    x = 0.62 + i * 4.12
    box(slide, x, 2.28, 3.83, 2.55, WHITE)
    text(slide, n, x + 0.27, 2.6, 0.75, 0.34, 17, accent, True)
    text(slide, heading, x + 0.27, 3.16, 3.2, 0.4, 20, INK, True)
    text(slide, body, x + 0.27, 3.84, 3.16, 0.61, 13, MUTED)
text(slide, "汇报结构：议题 → 受影响者 → 行动建议 → 证据", 0.88, 5.58, 11.6, 0.43, 17, NAVY, True, align=PP_ALIGN.CENTER)

# 13. Summary
slide = new_slide("本章小结", "回顾", 13,
                  "离场问题：写下你最想继续了解的一个 ESG 议题，以及一个理由。")
summary = [
    ("01", "ESG 看什么", "环境、社会与治理三个相互关联的维度", GREEN),
    ("02", "ESG 怎么用", "识别议题，管理行动，跟踪并沟通结果", BLUE),
    ("03", "ESG 与 SDGs", "目标议程与组织视角不同，可由具体行动衔接", ORANGE),
    ("04", "中国实践", "结合发展情境，把议题落到责任、决策与证据", "8A6CB1"),
]
for i, (n, head, body, accent) in enumerate(summary):
    y = 1.82 + i * 1.0
    text(slide, n, 0.9, y + 0.08, 0.65, 0.32, 15, accent, True)
    text(slide, head, 1.78, y + 0.05, 2.0, 0.37, 17, INK, True)
    text(slide, body, 4.12, y + 0.07, 7.7, 0.38, 14, MUTED)
    box(slide, 0.88, y + 0.66, 11.35, 0.018, LINE, radius=False)
box(slide, 0.88, 6.08, 11.35, 0.56, NAVY)
text(slide, "ESG不是“贴标签”，而是把重要议题变成可负责、可衡量、可改进的行动。", 1.12, 6.24, 10.85, 0.24, 14, WHITE, True, align=PP_ALIGN.CENTER)

# 14. Sources and scope
slide = new_slide("参考资料与使用说明", "拓展阅读", 14,
                  "本课件依据公开资料独立整理；未复现或声称引用教材原文。")
text(slide, "发展历程与全球框架", 0.72, 1.78, 5.4, 0.3, 14, GREEN, True)
for i, (label, url) in enumerate(SOURCES[:4]):
    add_linked_source(slide, f"• {label}", url, 0.82, 2.2 + i * 0.58, 11.6)
text(slide, "中国政策与披露", 0.72, 4.65, 5.4, 0.3, 14, BLUE, True)
for i, (label, url) in enumerate(SOURCES[4:]):
    add_linked_source(slide, f"• {label}", url, 0.82, 5.05 + i * 0.48, 11.6)
text(slide, "资料核验：2026-10-16  ·  公开资料用于课堂导读，具体规则适用范围以发布机构原文为准。", 0.74, 6.52, 11.9, 0.28, 9, MUTED)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {OUTPUT} ({len(prs.slides)} slides)")
